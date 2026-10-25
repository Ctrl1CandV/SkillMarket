#!/usr/bin/env python3
"""Rebuild eight academic chart samples; every value is synthetic demonstration data.

Requires python-pptx, matplotlib and numpy. Native charts contain embedded Excel
data; response surfaces/contours and boxplots are code-rendered PNG assets.
This sample builder is not a plotting tutorial or a general charting library.
"""
import argparse
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import sys

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.chart.datalabel import DataLabels
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_DATA_LABEL_POSITION,
                             XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DEMO = '演示数据，非研究结果'
BLUE, GRAY = '2764A5', '637B98'


def rgb(value):
    return RGBColor.from_string(value)


class Gallery:
    def __init__(self, font):
        catalog = json.loads((ROOT / 'assets' / 'themes.json').read_text(encoding='utf-8'))
        self.theme = deepcopy(next(t for t in catalog['themes'] if t['id'] == 'seminar-blue'))
        self.font = font
        self.prs = Presentation()
        self.prs.slide_width = Inches(catalog['canvas']['width_in'])
        self.prs.slide_height = Inches(catalog['canvas']['height_in'])
        self.prs.core_properties.title = '学术图表展示样张'
        self.prs.core_properties.subject = DEMO

    def style(self, run, size, ink=None, bold=False):
        run.font.name, run.font.size = self.font, Pt(size)
        run.font.bold, run.font.color.rgb = bold, rgb(ink or self.theme['text'])
        ea = OxmlElement('a:ea')
        ea.set('typeface', self.font)
        run._r.get_or_add_rPr().append(ea)

    def text(self, slide, value, x, y, w, h, size=18, ink=None, bold=False, center=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i, content in enumerate(value.split('\n')):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.space_after, p.line_spacing = Pt(5), 1.12
            if center:
                p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = content
            self.style(run, size, ink, bold)
        return box

    def page(self, title, kind, notes):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(self.theme['background'])
        self.text(slide, f'图表样张 / {kind}', .65, .32, 11.8, .32, 14, BLUE, True)
        self.text(slide, title, .65, .93, 12.0, .68, 30, bold=True)
        self.text(slide, DEMO, .65, 7.06, 11.2, .22, 11, self.theme['muted'])
        self.text(slide, f'{len(self.prs.slides):02d}', 12.12, 7.02, .5, .3, 12,
                  self.theme['muted'])
        slide.notes_slide.notes_text_frame.text = (
            f'{DEMO}。全部数值由本脚本明确设定或确定性生成，无外部数据来源。\n'
            f'主题：seminar-blue；用途：向 AI 展示可选择的学术图表与呈现效果。\n{notes}'
        )
        return slide

    def native(self, slide, kind, data, x, y, w, h, colors=(GRAY, BLUE), axes=True):
        chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        chart.has_title = False
        chart.font.name, chart.font.size, chart.font.color.rgb = self.font, Pt(18), rgb(self.theme['text'])
        chart.has_legend = len(chart.series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.name, chart.legend.font.size = self.font, Pt(18)
        for series, ink in zip(chart.series, colors):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = rgb(ink)
            series.format.line.color.rgb = rgb(ink)
            series.format.line.width = Pt(2.5)
        if axes:
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.name, axis.tick_labels.font.size = self.font, Pt(18)
                axis.tick_labels.font.color.rgb = rgb(self.theme['muted'])
                axis.major_tick_mark = XL_TICK_MARK.NONE
                axis.minor_tick_mark = XL_TICK_MARK.NONE
                axis.format.line.color.rgb = rgb(self.theme['rule'])
            chart.value_axis.has_major_gridlines = True
            grid = chart.value_axis.major_gridlines.format.line
            grid.color.rgb, grid.width = rgb('DCE4ED'), Pt(.7)
        return chart

    def axis_title(self, axis, value):
        axis.has_title = True
        frame = axis.axis_title.text_frame
        frame.clear()
        run = frame.paragraphs[0].add_run()
        run.text = value
        self.style(run, 18)

    def table(self, slide, rows, x, y, w, h, widths=None, fills=None, numeric_right=False):
        table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
        if widths:
            for col, fraction in zip(table.columns, widths):
                col.width = Inches(w * fraction)
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(.10)
                cell.margin_top = cell.margin_bottom = Inches(.06)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fills[i][j] if fills else
                                               (self.theme['surface'] if i == 0 else self.theme['background']))
                for edge in ('lnL', 'lnR', 'lnT', 'lnB'):
                    ln = OxmlElement('a:' + edge)
                    ln.append(OxmlElement('a:noFill'))
                    cell._tc.get_or_add_tcPr().append(ln)
                p = cell.text_frame.paragraphs[0]
                p.alignment = (PP_ALIGN.LEFT if j == 0 else
                               PP_ALIGN.RIGHT if numeric_right and i > 0 else PP_ALIGN.CENTER)
                run = p.add_run()
                run.text = str(value)
                self.style(run, 18, bold=i == 0)
        return table

    def picture(self, slide, figure, x, y, w, h, description):
        stream = BytesIO()
        figure.savefig(stream, format='png', dpi=300, facecolor='#' + self.theme['background'])
        plt.close(figure)
        stream.seek(0)
        shape = slide.shapes.add_picture(stream, Inches(x), Inches(y), Inches(w), Inches(h))
        shape._element.nvPicPr.cNvPr.set('descr', description)


def trend(g):
    a, b = [52, 62, 69, 73, 75, 76], [52, 63, 71, 77, 80, 82]
    s = g.page('趋势观察：两条曲线在后段逐渐分开', '折线趋势',
               f'原生折线图，含嵌入工作簿。检查点1–6；基线={a}；配置B={b}。\n'
               '呈现目的：连续横轴上的变化与局部差异；末点差6个百分点，不是相对提升6%。单条轨迹无误差估计。')
    g.text(s, '同一演示验证集；6 个检查点；每个配置展示一条轨迹', .7, 1.68, 11.9, .5)
    d = CategoryChartData(); d.categories = list(range(1, 7))
    d.add_series('基线', a); d.add_series('配置 B', b)
    c = g.native(s, XL_CHART_TYPE.LINE_MARKERS, d, .65, 2.18, 9.3, 4.55)
    c.value_axis.minimum_scale, c.value_axis.maximum_scale, c.value_axis.major_unit = 0, 100, 20
    g.axis_title(c.value_axis, '准确率 ↑ (%)'); g.axis_title(c.category_axis, '检查点')
    for series, ink, marker in zip(c.series, (GRAY, BLUE), (XL_MARKER_STYLE.SQUARE, XL_MARKER_STYLE.CIRCLE)):
        series.marker.style, series.marker.size = marker, 7
        series.marker.format.fill.solid(); series.marker.format.fill.fore_color.rgb = rgb(ink)
        series.marker.format.line.color.rgb = rgb(ink)
    c.series[0].format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    g.text(s, '末个检查点', 10.1, 3.05, 2.5, .4, 18, BLUE, True)
    g.text(s, '82% vs 76%\n相差 6 个百分点', 10.1, 3.64, 2.5, 1.3, 20, bold=True)
    g.text(s, '单条演示轨迹\n不能说明稳定提升', 10.1, 5.16, 2.5, .9, 18, g.theme['muted'])


def comparison(g):
    a, b = [68, 72, 75], [70, 74, 76]
    s = g.page('公平比较：固定预算后再比较配置差异', '分组柱状',
               f'原生分组柱状图，含嵌入数据。任务甲乙丙；基线={a}；配置B={b}。\n'
               '呈现目的：共享零点和相同尺度的组内比较；各组同一划分、训练预算和评测协议。每项为一次演示值，无误差条。')
    g.text(s, '组内控制：同一划分、同一训练预算、同一评测协议；每项为单次演示值', .7, 1.68, 11.9, .5)
    d = CategoryChartData(); d.categories = ['任务甲', '任务乙', '任务丙']
    d.add_series('基线', a); d.add_series('配置 B', b)
    c = g.native(s, XL_CHART_TYPE.COLUMN_CLUSTERED, d, .7, 2.25, 11.9, 4.45)
    c.value_axis.minimum_scale, c.value_axis.maximum_scale, c.value_axis.major_unit = 0, 100, 20
    g.axis_title(c.value_axis, '准确率 ↑ (%)')
    c.plots[0].gap_width = 95
    c.plots[0].overlap = -20
    c.plots[0].has_data_labels = True
    labels = c.plots[0].data_labels
    labels.show_legend_key = labels.show_category_name = labels.show_series_name = labels.show_percentage = False
    labels.show_value = True
    labels.position, labels.font.name, labels.font.size = XL_DATA_LABEL_POSITION.OUTSIDE_END, g.font, Pt(18)


def heatmap(g):
    values = np.array([[61, 65, 68, 70], [66, 72, 76, 74], [68, 75, 81, 78], [67, 73, 77, 75]])
    s = g.page('二维扫描：局部高值需要连同邻域一起看', '原生热力图',
               f'原生表格单元格与色标形状可编辑。行=候选数，列=温度；准确率(%)={values.tolist()}。\n'
               '呈现目的：离散超参数网格的局部结构；原值映射固定色标60–85%，无行列归一化。'
               '16格均有有效演示值，无缺失值或0。每格一次演示值，不能据此确定最优超参数。')
    g.text(s, '准确率 ↑ (%)；每格单次演示值；按原值映射固定色标 60–85%', .7, 1.68, 11.9, .5)
    rows = [['候选数 / 温度', '0.2', '0.4', '0.6', '0.8']]
    rows += [[str(k)] + list(v) for k, v in zip([10, 20, 30, 40], values)]
    def tint(v):
        t = (v - 60) / 25
        return ''.join(f'{round(a + (b-a)*t):02X}' for a, b in zip((239, 245, 251), (96, 155, 207)))
    fills = [[g.theme['surface']] * 5] + [[g.theme['surface']] + [tint(v) for v in row] for row in values]
    g.table(s, rows, .75, 2.48, 9.65, 3.48, [.26, .185, .185, .185, .185], fills)
    for i, v in enumerate(range(60, 86, 5)):
        shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.95), Inches(2.5+i*.48), Inches(.45), Inches(.5))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(tint(v)); shape.line.fill.background()
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        g.text(s, str(v), 11.57, 2.52+i*.48, .65, .35)
    g.text(s, '色标 (%)', 10.85, 5.65, 1.6, .4)
    g.text(s, '高值位于中部；边缘回落。保留完整网格，避免只摘出最高一格。', .75, 6.25, 11.9, .5, 20)


def tradeoff(g):
    points = [('A', 8, 53), ('B', 17, 68), ('C', 29, 79), ('D', 45, 83), ('E', 53, 84)]
    s = g.page('质量与成本：后段收益变小，成本仍在增加', '原生 XY 散点',
               f'原生XY散点，含嵌入工作簿。点为(配置, 延迟ms, 准确率%)：{points}。\n'
               '呈现目的：联合读取两个数值轴；各点为同一协议下的演示配置，不代表误差分布或实际Pareto结论。')
    g.text(s, '同一演示评测协议；每个点代表一个配置；横轴越左越快，纵轴越高越准', .7, 1.68, 11.9, .5)
    d = XyChartData()
    for name, x, y in points:
        d.add_series(name).add_data_point(x, y)
    c = g.native(s, XL_CHART_TYPE.XY_SCATTER, d, .7, 2.25, 11.9, 4.48, (GRAY, GRAY, BLUE, GRAY, GRAY))
    c.has_legend = False
    for axis, limit, step in [(c.category_axis, 60, 10), (c.value_axis, 100, 20)]:
        axis.minimum_scale, axis.maximum_scale, axis.major_unit = 0, limit, step
    g.axis_title(c.category_axis, '推理延迟 ↓ (ms)'); g.axis_title(c.value_axis, '准确率 ↑ (%)')
    for series, ink in zip(c.series, (GRAY, GRAY, BLUE, GRAY, GRAY)):
        series.marker.style, series.marker.size = XL_MARKER_STYLE.CIRCLE, 12
        series.marker.format.fill.solid(); series.marker.format.fill.fore_color.rgb = rgb(ink)
        series.marker.format.line.color.rgb = rgb(ink)
    # python-pptx lacks ScatterPlot.has_data_labels; insert the native OOXML node.
    node = OxmlElement('c:dLbls')
    plot = c.plots[0]._element
    plot.insert(list(plot).index(plot.ser_lst[-1]) + 1, node)
    labels = DataLabels(node)
    labels.show_legend_key = labels.show_category_name = labels.show_percentage = False
    labels.show_value, labels.show_series_name = False, True
    labels.position = XL_DATA_LABEL_POSITION.ABOVE
    labels.font.name, labels.font.size = g.font, Pt(18)


def response(g):
    s = g.page('响应面与等高线：同一函数的两种空间读法', '三维与二维对照',
               '可重建PNG绘图资产，不是原生可编辑曲面/等高线。公式 z=80-20(x-0.6)^2-30(y-0.4)^2；'
               'x,y=np.linspace(0,1,41)，meshgrid后确定性求值。无测量样本或误差。\n'
               '呈现目的：3D显示整体形态，等高线便于定位峰值与比较水平集；两图共享相同网格、尺度与色标。')
    g.text(s, '左：空间形态；右：峰值位置与水平集。两图共享函数、网格和色标。', .7, 1.68, 11.9, .5)
    x, y = np.meshgrid(np.linspace(0, 1, 41), np.linspace(0, 1, 41))
    z = 80 - 20*(x-.6)**2 - 30*(y-.4)**2
    fig = plt.figure(figsize=(12, 4.4))
    ax = fig.add_axes([.01, .08, .45, .87], projection='3d')
    ax.plot_surface(x, y, z, cmap='Blues', vmin=60, vmax=80, linewidth=0, antialiased=True)
    ax.view_init(elev=27, azim=-58); ax.set_box_aspect((1, 1, .65))
    ax.set(xlabel='x', ylabel='y', zlabel='z', xticks=[0, .5, 1], yticks=[0, .5, 1], zticks=[60, 70, 80], zlim=(60, 80))
    ax.set_xlabel('x', labelpad=14); ax.set_ylabel('y', labelpad=14); ax.set_zlabel('z', labelpad=10)
    ax2 = fig.add_axes([.56, .19, .31, .72])
    contours = ax2.contourf(x, y, z, levels=np.linspace(60, 80, 9), cmap='Blues', vmin=60, vmax=80)
    ax2.contour(x, y, z, levels=[65, 70, 75], colors='#637B98', linewidths=.8)
    ax2.set(xlabel='x', ylabel='y', xticks=[0, .5, 1], yticks=[0, .5, 1], aspect='equal')
    fig.colorbar(contours, cax=fig.add_axes([.91, .23, .022, .62]), ticks=[60, 70, 80], label='z')
    g.picture(s, fig, .65, 2.12, 12.0, 4.4, '确定性函数响应面及等高线；PNG绘图资产，公式与41×41网格定义见备注及源脚本。')
    g.text(s, 'z = 80 − 20(x − 0.6)² − 30(y − 0.4)²；0 ≤ x, y ≤ 1', .75, 6.35, 11.9, .4, 18)
    g.text(s, '3D 与等高线为可重建 PNG 绘图资产', .75, 6.76, 11.9, .27, 14, g.theme['muted'])


def proportions(g):
    s = g.page('组成与比较：环形看整体，条形对齐比例', '部分与整体',
               '两个原生图表均含嵌入数据。配置A=[60,25,15]、配置B=[48,32,20]；类别为计算、检索、其他，单位%。\n'
               '呈现目的：环形图表示单个整体；100%堆叠条形在一致顺序下比较组成。每组总计100%，无样本量或统计误差。')
    g.text(s, '同一耗时口径；类别顺序与颜色一致；每组总和均为 100%', .7, 1.68, 11.9, .5)
    parts, a, b = ['计算', '检索', '其他'], [60, 25, 15], [48, 32, 20]
    colors = (BLUE, GRAY, 'BBC9D8')
    d = CategoryChartData(); d.categories = parts; d.add_series('配置 A', a)
    c = g.native(s, XL_CHART_TYPE.DOUGHNUT, d, .75, 2.65, 3.4, 3.72, axes=False)
    c.plots[0]._element.xpath('./c:holeSize')[0].set('val', '63')
    for point, ink in zip(c.series[0].points, colors):
        point.format.fill.solid(); point.format.fill.fore_color.rgb = rgb(ink)
        point.format.line.fill.background()
    # Outside native text avoids label collisions and PowerPoint's doughnut dLblPos incompatibility.
    c.plots[0].has_data_labels = False
    g.text(s, '配置 A', 1.60, 4.28, 1.7, .5, 20, bold=True, center=True)
    for i, (name, value, ink) in enumerate(zip(parts, a, colors)):
        y = 3.32 + i*.75
        swatch = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.40), Inches(y+.1), Inches(.16), Inches(.16))
        swatch.fill.solid(); swatch.fill.fore_color.rgb = rgb(ink); swatch.line.fill.background()
        swatch._element.spPr.append(OxmlElement('a:effectLst'))
        g.text(s, f'{name} {value}%', 4.68, y, 1.4, .5, 18)
    d = CategoryChartData(); d.categories = ['配置 B', '配置 A']  # Horizontal categories display bottom to top.
    for name, av, bv in zip(parts, a, b):
        d.add_series(name, [bv, av])
    c = g.native(s, XL_CHART_TYPE.BAR_STACKED_100, d, 6.2, 2.65, 6.4, 3.72, colors)
    c.value_axis.tick_labels.number_format = '0%'
    c.plots[0].gap_width = 80
    c.plots[0].has_data_labels = True
    labels = c.plots[0].data_labels
    labels.position, labels.font.name, labels.font.size = XL_DATA_LABEL_POSITION.CENTER, g.font, Pt(18)
    labels.number_format = '0"%"'
    for series, ink in zip(c.series, ['FFFFFF', 'FFFFFF', g.theme['text']]):
        labels = series.data_labels
        labels.show_value = True
        labels.show_legend_key = labels.show_category_name = labels.show_series_name = labels.show_percentage = False
        labels.position, labels.number_format = XL_DATA_LABEL_POSITION.CENTER, '0"%"'
        labels.font.name, labels.font.size, labels.font.color.rgb = g.font, Pt(18), rgb(ink)
    g.text(s, '计算占比从 60% 变为 48%：减少 12 个百分点。', .75, 6.47, 11.9, .45, 20)


def distribution(g):
    a = [16, 17, 17, 18, 18, 19, 19, 20, 20, 21, 22, 29]
    b = [15, 16, 16, 17, 17, 18, 18, 18, 19, 19, 20, 23]
    s = g.page('分布观察：箱线保留概括，原始点保留样本', '箱线与原始点',
               f'PNG绘图资产；可从源脚本重建。配置A={a}；配置B={b}；每组n=12个演示请求，单位ms。\n'
               '箱体=第25–75百分位，横线=中位数，须=1.5×IQR内最远样本；全部原始点均显示。'
               '横向偏移为固定的12个位置，只用于避让重叠，不代表第三变量。'
               '呈现目的：比较集中、离散与尾部，不做显著性推断。')
    g.text(s, '每组 n = 12 个演示请求；箱体为四分位区间；全部原始点均显示', .7, 1.68, 11.9, .5)
    fig, ax = plt.subplots(figsize=(9.3, 4.6))
    fig.subplots_adjust(left=.13, right=.97, bottom=.18, top=.95)
    boxes = ax.boxplot([a, b], positions=[1, 2], widths=.38, patch_artist=True, showfliers=False,
                       medianprops={'color': '#182B42', 'linewidth': 2}, whiskerprops={'linewidth': 1.5}, capprops={'linewidth': 1.5})
    for i, (box, values, ink) in enumerate(zip(boxes['boxes'], [a, b], [GRAY, BLUE]), 1):
        box.set(facecolor='#'+ink, alpha=.22, edgecolor='#'+ink)
        jitter = np.array([-.10, .07, -.03, .11, -.06, .02, -.12, .09, -.01, .05, -.08, 0])
        ax.scatter(i+jitter, values, s=60, color='#'+ink, edgecolors='white', linewidths=.7, zorder=3)
    ax.set(xticks=[1, 2], xticklabels=['配置 A', '配置 B'], ylabel='请求延迟 ↓ (ms)', ylim=(10, 32), yticks=[10, 15, 20, 25, 30])
    ax.grid(axis='y', color='#DCE4ED'); ax.set_axisbelow(True)
    ax.spines[['top', 'right']].set_visible(False)
    g.picture(s, fig, .65, 2.17, 9.3, 4.6, '箱线图叠加全部12个原始点；PNG资产；原始数值与统计定义见备注及源脚本。')
    g.text(s, '尾部样本仍可见', 10.08, 3.12, 2.55, .5, 18, BLUE, True)
    g.text(s, '不只看均值\n也看离散与长尾', 10.08, 3.76, 2.55, 1.1, 20, bold=True)
    g.text(s, '须：1.5 × IQR 内\n最远的实际样本', 10.08, 5.2, 2.55, .9, 18, g.theme['muted'])


def result_table(g):
    runs = {'基线': ([71, 72, 73], [18, 19, 20]), '配置 B': ([73, 74, 75], [21, 22, 23]),
            '配置 C': ([72, 73, 74], [19, 20, 21])}
    rows = [['配置', '准确率 ↑ (%)', 'P95 延迟 ↓ (ms)', '运行次数 n']]
    for name, (accuracy, latency) in runs.items():
        rows.append([name, f'{np.mean(accuracy):.1f} ± {np.std(accuracy, ddof=1):.1f}',
                     f'{np.mean(latency):.1f} ± {np.std(latency, ddof=1):.1f}', '3'])
    s = g.page('结果表：方向、单位与误差含义一起呈现', '原生科研结果表',
               f'原生可编辑表格。各配置三次演示运行的(准确率%,P95延迟ms)={runs}。\n'
               '均值与样本标准差由这3个值计算，ddof=1；n是运行次数，不是请求数。所有配置同划分、同训练预算和同评测协议。'
               '呈现目的：多指标权衡；仅加粗配置B的最高准确率，未进行显著性检验，不能把差异当作稳定结论。')
    g.text(s, '同一划分、同一训练预算与同一评测协议；n = 每配置的演示运行次数', .7, 1.68, 11.9, .5)
    table = g.table(s, rows, .75, 2.66, 11.8, 2.8, [.20, .29, .33, .18], numeric_right=True)
    for run in table.cell(2, 1).text_frame.paragraphs[0].runs:
        run.font.bold = True; run.font.color.rgb = rgb(BLUE)
    g.text(s, '数值为 mean ± SD；SD 为 3 次运行的样本标准差（ddof = 1）。', .75, 5.82, 11.8, .5, 18)
    g.text(s, '配置 B 的准确率较高，同时延迟增加；此样张未做显著性检验。', .75, 6.43, 11.8, .5, 20)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True, help='New PPTX destination; refuses to overwrite')
    parser.add_argument('--font', default='Microsoft YaHei', help='Installed font for Chinese and chart labels')
    args = parser.parse_args()
    try:
        if args.output.exists():
            raise ValueError('Output already exists; choose a new filename')
        font_manager.findfont(args.font, fallback_to_default=False)
        plt.rcParams.update({'font.family': args.font, 'font.size': 18, 'axes.labelsize': 18,
                             'xtick.labelsize': 18, 'ytick.labelsize': 18, 'axes.unicode_minus': False,
                             'text.color': '#182B42', 'axes.labelcolor': '#182B42',
                             'axes.facecolor': '#F7F9FC', 'figure.facecolor': '#F7F9FC'})
        g = Gallery(args.font)
        for draw in (trend, comparison, heatmap, tradeoff, response, proportions, distribution, result_table):
            draw(g)
        # Chart APIs set Latin fonts only; explicitly supply the same East Asian font.
        for slide in g.prs.slides:
            for shape in slide.shapes:
                if shape.has_chart:
                    for props in shape.chart._chartSpace.xpath('.//a:defRPr | .//a:rPr | .//a:endParaRPr'):
                        ea = props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
                        if ea is None:
                            ea = OxmlElement('a:ea'); props.append(ea)
                        ea.set('typeface', g.font)
        stream = BytesIO(); g.prs.save(stream)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open('xb') as destination:
            destination.write(stream.getvalue())
        print(json.dumps({'output': str(args.output), 'slides': 8, 'synthetic_content': True,
                          'native_charts': 5, 'png_assets': 2, 'visually_reviewed': False}, ensure_ascii=False))
        return 0
    except (OSError, ValueError, KeyError) as exc:
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
