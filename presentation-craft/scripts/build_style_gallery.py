#!/usr/bin/env python3
"""Build editable academic style samples with explicitly synthetic content.

Requires python-pptx. Does not render slides or certify their visual quality.
"""
import argparse
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
METHOD_TITLE = '固定候选预算后，重排模块如何改变候选排序？'
RESULT_TITLE = '验证集差异很小，仍需重复实验确认'
STEPS = [('查询 q', '输入问题'), ('候选集', '固定 K = 50'),
         ('重排模块', '本轮唯一改动'), ('输出排序', '同一评测协议')]
ROWS = [('基线', '11', '71.2', '18.4'), ('基线', '29', '71.8', '18.5'),
        ('加重排', '11', '71.5', '21.6'), ('加重排', '29', '71.6', '21.7')]


def color(value):
    return RGBColor.from_string(value)


def text(slide, theme, value, x, y, w, h, size=None, ink=None, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, content in enumerate(value.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = 1.12
        paragraph.space_after = Pt(6)
        run = paragraph.add_run()
        run.text = content
        style_run(run, theme, size or theme['sizes']['body'], ink or theme['text'], bold)
    return box


def style_run(run, theme, size, ink, bold=False):
    run.font.name = theme['fonts']['latin']
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(ink)
    properties = run._r.get_or_add_rPr()
    east_asian = properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    if east_asian is None:
        east_asian = OxmlElement('a:ea')
        properties.append(east_asian)
    east_asian.set('typeface', theme['fonts']['cjk'])


def rect(slide, x, y, w, h, fill, edge=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if edge:
        shape.line.color.rgb = color(edge)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    # A stock Office rectangle may inherit a theme shadow despite our fill.
    # An explicit empty effect list keeps the academic diagram flat.
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape


def line(slide, x1, y1, x2, y2, ink, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color(ink)
    shape.line.width = Pt(1.2)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    if arrow:
        end = OxmlElement('a:tailEnd')
        end.set('type', 'triangle')
        shape.line._get_or_add_ln().append(end)
    return shape


def header(slide, theme, index, page, title, kind):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(theme['background'])
    text(slide, theme, f'{index:02d}  {theme["name"]}  /  {kind}', .65, .32, 10.7, .3,
         14, theme['accent'], True)
    text(slide, theme, title, .65, .94, 12.0, 1.08, theme['sizes']['title'], bold=True)
    text(slide, theme, '样式预览 · 演示数据与示意流程，非研究结论', .65, 7.06, 11.2, .22,
         theme['sizes']['source'], theme['muted'])
    text(slide, theme, f'{page:02d}', 12.15, 7.03, .5, .25, 12, theme['muted'])


def pipeline(slide, theme, x, y, w, vertical=False, open_nodes=False, focus=2):
    gap = .33 if not vertical else .23
    node_w = w if vertical else (w - gap * 3) / 4
    node_h = .78 if vertical else 1.18
    for i, (name, detail) in enumerate(STEPS):
        xx = x if vertical else x + i * (node_w + gap)
        yy = y + i * (node_h + gap) if vertical else y
        if not open_nodes:
            rect(slide, xx, yy, node_w, node_h, theme['surface'], theme['accent'] if i == focus else theme['rule'])
        text(slide, theme, name, xx + .12, yy + .12, node_w - .24, .4,
             20 if node_w > 2 else 18, theme['accent'] if i == focus else theme['text'], i == focus)
        text(slide, theme, detail, xx + .12, yy + (.48 if vertical else .69), node_w - .24, .24,
             13 if vertical else 14, theme['muted'])
        if i < 3:
            if vertical:
                line(slide, xx + node_w / 2, yy + node_h + .02,
                     xx + node_w / 2, yy + node_h + gap - .02, theme['accent'], True)
            else:
                line(slide, xx + node_w + .03, yy + .55,
                     xx + node_w + gap - .03, yy + .55, theme['accent'], True)


def explanation(slide, theme, x, y, w, label='研究问题', detail='比较候选内排序，\n不扩大候选集。'):
    text(slide, theme, label, x, y, w, .36, 16, theme['accent'], True)
    text(slide, theme, detail, x, y + .55, w, 1.4, 24, bold=True)
    text(slide, theme, '控制条件\n同一数据划分与候选数', x, y + 2.2, w, 1.0, 18, theme['muted'])


def method_page(slide, theme):
    layout = theme['layout']
    if layout == 'clean_split':
        explanation(slide, theme, .65, 2.25, 3.35)
        pipeline(slide, theme, 4.45, 2.7, 8.2)
        text(slide, theme, '只改变评分与排序；候选集保持不变。', 4.45, 4.5, 8.2, .85, 22)
    elif layout == 'paper_plate':
        text(slide, theme, '方法示意｜从输入到输出，沿箭头阅读', .7, 2.2, 11.9, .45, 18, theme['muted'])
        pipeline(slide, theme, .7, 3.05, 11.9, open_nodes=True)
        line(slide, .7, 4.62, 12.6, 4.62, theme['rule'])
        text(slide, theme, '研究问题：固定候选数后，排序变化是否值得额外延迟？', .7, 5.0, 11.9, .8, 23)
    elif layout == 'lab_notebook':
        explanation(slide, theme, .65, 2.25, 3.3, '本轮改动', '只加入重排，\n其余设置保持不变。')
        pipeline(slide, theme, 4.5, 2.35, 8.1)
        for y, key, value in [(4.25, '保持', '查询、候选集与评测脚本'), (5.23, '观察', '排序效果与额外推理延迟')]:
            line(slide, 4.5, y, 12.6, y, theme['rule'])
            text(slide, theme, key, 4.5, y + .18, 1.0, .5, 18, theme['accent'], True)
            text(slide, theme, value, 5.7, y + .18, 6.8, .65, 21)
    elif layout == 'method_bands':
        text(slide, theme, '输入', .75, 2.2, 2.2, .4, 18, theme['muted'])
        text(slide, theme, '候选内变换', 4.0, 2.2, 5.0, .4, 18, theme['accent'], True)
        text(slide, theme, '输出', 10.0, 2.2, 2.2, .4, 18, theme['muted'])
        rect(slide, 3.6, 2.95, 6.05, 1.8, theme['surface'])
        pipeline(slide, theme, .75, 3.25, 11.8, open_nodes=True)
        text(slide, theme, '局部解释：重排模块接收固定候选集，仅更新候选的顺序。', .75, 5.3, 11.8, .9, 23)
    elif layout == 'result_focus':
        text(slide, theme, '比较必须先控制候选预算', .7, 2.2, 11.9, .6, 28, theme['accent'], True)
        text(slide, theme, '固定 K = 50 与数据划分，分别观察效果和延迟。', .7, 3.05, 11.9, .7, 22)
        pipeline(slide, theme, .7, 4.35, 11.9)
    elif layout == 'theorem':
        text(slide, theme, '排序操作的作用域', .7, 2.2, 7.7, .5, 18, theme['muted'])
        text(slide, theme, 'π = argsort sθ(q, d)', .9, 3.05, 7.0, .9, 32, theme['accent'], True)
        text(slide, theme, 'd ∈ Cₖ(q)，候选集不变', .9, 4.05, 7.0, .6, 24)
        text(slide, theme, '只更新顺序，不改变候选来源。', .9, 5.2, 7.0, .7, 22)
        line(slide, 8.65, 2.3, 8.65, 6.15, theme['rule'])
        text(slide, theme, 'q：输入查询\nd：候选文档\nK = 50：候选数\nsθ：重排评分', 9.0, 2.65, 3.5, 2.8, 21)
        text(slide, theme, '符号仅用于方法示意', 9.0, 5.65, 3.5, .5, 14, theme['muted'])
    elif layout == 'dark_system':
        rect(slide, .8, 2.35, 11.7, 2.8, theme['surface'])
        text(slide, theme, '检索与重排路径', 1.05, 2.58, 10.8, .4, 18, theme['accent'], True)
        pipeline(slide, theme, 1.05, 3.4, 11.2, open_nodes=True)
        text(slide, theme, '当前关注重排模块；其他模块是控制条件。', .8, 5.7, 11.7, .7, 23)
    elif layout == 'image_strip':
        text(slide, theme, '同一候选、相同尺度；图形只表示候选身份，不是实验图。',
             .7, 2.15, 11.9, .6, 18, theme['muted'])
        patterns = {'A': [(0, 0), (1, 0), (2, 0)], 'B': [(0, 0), (1, 1), (2, 0)],
                    'C': [(0, 1), (1, 0), (2, 1)], 'D': [(0, 0), (1, 0), (2, 1)]}
        for y, label, ordering in [(3.1, '重排前', 'ABCD'), (4.8, '重排后', 'BACD')]:
            text(slide, theme, label, .7, y + .35, 1.4, .5, 20, bold=True)
            for position, item in enumerate(ordering):
                x = 2.4 + position * 2.6
                rect(slide, x, y, 2.25, 1.25, theme['surface'], theme['accent'] if item == 'B' else theme['rule'])
                for xx, yy in patterns[item]:
                    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + .48 + xx * .48),
                                                Inches(y + .16 + yy * .18), Inches(.16), Inches(.16))
                    dot.fill.solid()
                    dot.fill.fore_color.rgb = color(theme['accent'] if item == 'B' else theme['muted'])
                    dot.line.fill.background()
                    dot._element.spPr.append(OxmlElement('a:effectLst'))
                text(slide, theme, f'{position + 1} / 候选 {item}', x + .25, y + .73, 1.9, .4, 18,
                     theme['accent'] if item == 'B' else theme['text'], item == 'B')
        text(slide, theme, '示意：候选 B 的位置前移；是否改善检索质量仍需结果验证。',
             .7, 6.4, 11.9, .5, 19, theme['accent'])
    elif layout == 'observation':
        text(slide, theme, '处理条件：固定候选集，只改变候选内排序', .7, 2.2, 11.9, .7, 23, bold=True)
        pipeline(slide, theme, .7, 3.4, 11.9, open_nodes=True)
        for x, label, detail in [( .7, '观察效果', '同一验证集上的检索指标'), (7.0, '观察代价', '同一协议下的推理延迟')]:
            line(slide, x, 5.12, x + 5.6, 5.12, theme['rule'])
            text(slide, theme, label, x, 5.35, 5.6, .4, 18, theme['accent'], True)
            text(slide, theme, detail, x, 5.92, 5.6, .6, 21)
    elif layout == 'editorial':
        explanation(slide, theme, .7, 2.2, 4.25, '方法要回答的问题', '额外排序成本，\n换来了什么？')
        pipeline(slide, theme, 6.2, 2.2, 5.9, vertical=True)
    elif layout == 'teaching':
        rect(slide, .7, 2.25, 11.9, .8, theme['surface'])
        text(slide, theme, '贯穿示例：同一个问题，比较重排前后的候选顺序。', .95, 2.43, 11.35, .5, 22)
        pipeline(slide, theme, .7, 3.8, 11.9)
        text(slide, theme, '解释时保留输入，逐步加入“候选集”和“评分”的含义。', .7, 5.65, 11.9, .7, 21)
    else:  # formal
        text(slide, theme, '研究问题', .7, 2.25, 2.1, .4, 17, theme['accent'], True)
        text(slide, theme, '固定预算下的排序质量与效率', 3.0, 2.2, 9.6, .7, 25, bold=True)
        pipeline(slide, theme, .7, 3.55, 11.9)
        text(slide, theme, '验证路径', .7, 5.4, 2.1, .4, 17, theme['accent'], True)
        text(slide, theme, '保持评测条件一致，再比较效果与延迟。', 3.0, 5.35, 9.6, .85, 22)


def result_table(slide, theme, x, y, w):
    dark = theme['layout'] == 'dark_system'
    ink = '182B42' if dark else theme['text']
    accent = '166B5B' if dark else theme['accent']
    surface = 'EAF1F5' if dark else theme['surface']
    if dark:
        rect(slide, x - .12, y - .15, w + .24, 3.6, 'FFFFFF')
    table = slide.shapes.add_table(5, 4, Inches(x), Inches(y), Inches(w), Inches(3.25)).table
    widths = [.25, .12, .34, .29]
    for column, fraction in zip(table.columns, widths):
        column.width = Inches(w * fraction)
    headings = ['方法', 'Seed', 'Recall@10 ↑ (%)', 'P95 ↓ (ms)']
    for row_index, values in enumerate([headings] + ROWS):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.margin_left = Inches(.13)
            cell.margin_right = Inches(.05)
            cell.margin_top = Inches(.13)
            cell.margin_bottom = Inches(.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(surface if row_index == 0 else ('FFFFFF' if dark else theme['background']))
            # Quiet native table borders; horizontal rules express row groups.
            props = cell._tc.get_or_add_tcPr()
            for edge in ('lnL', 'lnR', 'lnT', 'lnB'):
                border = OxmlElement('a:' + edge)
                border.append(OxmlElement('a:noFill'))
                props.append(border)
            paragraph = cell.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = value
            style_run(run, theme, 17 if row_index == 0 else 20,
                      accent if col_index == 0 and row_index > 2 else ink, row_index == 0)
    line(slide, x, y + .65, x + w, y + .65, theme['rule'])
    line(slide, x, y + 1.95, x + w, y + 1.95, theme['rule'])


def result_page(slide, theme):
    layout = theme['layout']
    side = layout in ('clean_split', 'lab_notebook', 'theorem', 'editorial', 'formal')
    text(slide, theme, '同一验证集与 K = 50；每种配置 2 次运行；以下均为演示数据。',
         .7, 2.1, 11.9, .6, 17, theme['muted'])
    if side:
        if layout == 'lab_notebook':
            explanation(slide, theme, .7, 3.0, 3.05, '观察与待解释', '效果方向不一致，\n延迟有所增加。')
            result_table(slide, theme, 4.2, 3.0, 8.4)
        else:
            result_table(slide, theme, .7, 3.0, 8.2)
            text(slide, theme, '当前观察', 9.35, 3.05, 3.25, .4, 17, theme['accent'], True)
            text(slide, theme, '效果方向\n并不一致', 9.35, 3.7, 3.25, 1.25, 25, bold=True)
            text(slide, theme, '需要增加重复，\n再判断是否保留。', 9.35, 5.1, 3.25, 1.1, 20)
    else:
        result_table(slide, theme, .8, 2.95, 11.7)
        text(slide, theme, '下一步：增加重复，核对波动与延迟，再决定是否保留重排。',
             .8, 6.45, 11.7, .45, 19, theme['accent'], True)


def build(catalog, output, selected=None, font=None):
    themes = catalog['themes']
    if selected:
        unknown = set(selected) - {theme['id'] for theme in themes}
        if unknown:
            raise ValueError('Unknown theme IDs: ' + ', '.join(sorted(unknown)))
        themes = [theme for theme in themes if theme['id'] in selected]
    presentation = Presentation()
    presentation.slide_width = Inches(catalog['canvas']['width_in'])
    presentation.slide_height = Inches(catalog['canvas']['height_in'])
    presentation.core_properties.title = '研究生汇报样式样张 · 方法与结果'
    presentation.core_properties.subject = '演示数据与示意流程，不是研究结论'
    for theme in themes:
        if font:
            theme['fonts']['cjk'] = font
        index = next(i for i, item in enumerate(catalog['themes'], 1) if item['id'] == theme['id'])
        for title, kind, draw in [(METHOD_TITLE, '方法示意', method_page), (RESULT_TITLE, '结果对照', result_page)]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            header(slide, theme, index, len(presentation.slides), title, kind)
            draw(slide, theme)
            slide.notes_slide.notes_text_frame.text = (
                f'样式：{theme["id"]} / {theme["name"]}。适用：{"、".join(theme["best_for"])}。\n'
                f'构图约定：{theme["recipe"]}\n'
                '本页全部数值与流程仅用于比较样式，非论文数据或实际实验。正式稿必须替换内容和备注。\n'
                '讲解提示：先说明比较问题和控制条件，再沿证据阅读；两次运行不能据此声称稳定提升。'
            )
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return len(presentation.slides)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True, help='Destination PPTX; refuses to overwrite')
    parser.add_argument('--theme', action='append', help='Theme ID; repeat to compare a subset (catalog order)')
    parser.add_argument('--themes-file', type=Path, default=ROOT / 'assets' / 'themes.json')
    parser.add_argument('--font', help='Installed CJK font override; verify on the presentation computer')
    args = parser.parse_args()
    try:
        if args.output.exists():
            raise ValueError('Output already exists; choose a new filename')
        catalog = json.loads(args.themes_file.read_text(encoding='utf-8'))
        pages = build(catalog, args.output, args.theme, args.font)
        print(json.dumps({'output': str(args.output.resolve()), 'slides': pages,
                          'synthetic_content': True, 'visually_reviewed': False}, ensure_ascii=False))
        return 0
    except (OSError, ValueError, KeyError) as exc:
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
